#!/usr/bin/env python3
import io
import json
import re
import sys
import time
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

SSE_URL = "https://query.sse.com.cn/security/stock/downloadStockListFile.do?csrcCode=&stockCode=&areaName=&stockType=1"
SZSE_URL = "https://www.szse.cn/api/report/ShowReport?SHOWTYPE=xlsx&CATALOGID=1110&TABKEY=tab1"

UA = "Mozilla/5.0 stock-rador-audit-universe/0.6"
MAIN_SH = ("600", "601", "603", "605")
MAIN_SZ = ("000", "001", "002", "003")
CHINEXT = ("300", "301")
ALLOWED_PREFIXES = MAIN_SH + MAIN_SZ + CHINEXT

def http_bytes(url, referer):
    last = None
    for attempt in range(1, 6):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA, "Referer": referer})
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = resp.read()
                if len(data) < 1000:
                    raise RuntimeError(f"response too small: {len(data)} bytes")
                return data, attempt
        except Exception as exc:
            last = exc
            if attempt == 5:
                break
            time.sleep(0.5 * (2 ** (attempt - 1)))
    raise RuntimeError(f"failed after retries: {url}: {last}")

def board_for(code):
    if code.startswith(MAIN_SH):
        return "sh_main"
    if code.startswith(MAIN_SZ):
        return "sz_main"
    if code.startswith(CHINEXT):
        return "chinext"
    if code.startswith("688"):
        return "star"
    return "other"

def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()

def code_from_value(value):
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        text = f"{value:06d}"
    elif isinstance(value, float) and value.is_integer():
        text = f"{int(value):06d}"
    else:
        text = clean_text(value)
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2].zfill(6)
    return text if re.fullmatch(r"\d{6}", text or "") else None

def likely_name(value):
    text = clean_text(value)
    if not text or re.fullmatch(r"\d+(?:\.\d+)?", text):
        return False
    if re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", text):
        return False
    lowered = text.lower()
    if any(word in lowered for word in ["代码", "简称", "名称", "日期", "行业", "板块", "公司"]):
        return False
    return len(text) <= 40

def extract_rows(xlsx_bytes, exchange):
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    candidates = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values = list(row)
            for idx, value in enumerate(values):
                code = code_from_value(value)
                if not code:
                    continue
                board = board_for(code)
                if exchange == "SH" and board not in ("sh_main", "star"):
                    continue
                if exchange == "SZ" and board not in ("sz_main", "chinext"):
                    continue
                name = ""
                # Official stock-list sheets place the security name close to the code.
                for j in range(idx + 1, min(len(values), idx + 5)):
                    if likely_name(values[j]):
                        name = clean_text(values[j])
                        break
                if not name:
                    for j in range(max(0, idx - 3), idx):
                        if likely_name(values[j]):
                            name = clean_text(values[j])
                            break
                if not name:
                    continue
                candidates[code] = {
                    "code": code,
                    "name": name,
                    "exchange": exchange,
                    "marketCode": 1 if exchange == "SH" else 0,
                    "board": board,
                }
    return list(candidates.values())

def exclusion_reason(item):
    code = item["code"]
    name = item["name"]
    board = item["board"]
    if board == "star":
        return "star_board"
    if board not in ("sh_main", "sz_main", "chinext"):
        return "outside_frozen_boards"
    if re.search(r"\*?ST", name, re.I):
        return "st"
    if "退" in name:
        return "delisting_marker"
    if re.match(r"^[NC]", name, re.I):
        return "recent_listing_marker"
    return None

def ticker(item):
    return f'{item["code"]}.{"SH" if item["exchange"] == "SH" else "SZ"}'

def main():
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python scripts/fetch-a-share-audit-universe-official.py <output.json> [as-of-date]")
    output = Path(sys.argv[1])
    as_of = sys.argv[2] if len(sys.argv) > 2 else datetime.now(timezone.utc).date().isoformat()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", as_of):
        raise RuntimeError("as-of-date must be YYYY-MM-DD")
    if output.exists():
        raise RuntimeError(f"refusing to overwrite existing universe {output}")

    sse_bytes, sse_attempts = http_bytes(SSE_URL, "https://www.sse.com.cn/assortment/stock/list/share/")
    szse_bytes, szse_attempts = http_bytes(SZSE_URL, "https://www.szse.cn/market/product/stock/list/index.html")

    sse = extract_rows(sse_bytes, "SH")
    szse = extract_rows(szse_bytes, "SZ")
    if len(sse) < 1000:
        raise RuntimeError(f"unexpected SSE parsed count: {len(sse)}")
    if len(szse) < 2000:
        raise RuntimeError(f"unexpected SZSE parsed count: {len(szse)}")

    all_items = sse + szse
    seen = set()
    included, excluded = [], []
    for item in all_items:
        t = ticker(item)
        if t in seen:
            raise RuntimeError(f"duplicate ticker across official files: {t}")
        seen.add(t)
        row = {**item, "ticker": t}
        reason = exclusion_reason(item)
        if reason:
            excluded.append({**row, "exclusionReason": reason})
        else:
            included.append(row)

    included.sort(key=lambda x: x["ticker"])
    excluded.sort(key=lambda x: x["ticker"])
    counts = {}
    for item in excluded:
        counts[item["exclusionReason"]] = counts.get(item["exclusionReason"], 0) + 1

    now = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    universe = {
        "schemaVersion": "1.0",
        "universeVersion": "cn-a-share-main-chinext-official-v1",
        "generatedAt": now,
        "asOfDate": as_of,
        "source": {
            "provider": "SSE + SZSE official stock-list downloads",
            "retrievedAt": now,
            "files": [
                {"exchange": "SH", "url": SSE_URL, "parsedCount": len(sse), "attempts": sse_attempts},
                {"exchange": "SZ", "url": SZSE_URL, "parsedCount": len(szse), "attempts": szse_attempts},
            ],
        },
        "inclusionRule": {
            "exchanges": ["Shanghai", "Shenzhen"],
            "boards": ["Shanghai main board", "Shenzhen main board", "ChiNext"],
            "includePrefixes": list(ALLOWED_PREFIXES),
            "exclusions": [
                "STAR Market (688xxx)",
                "ST/*ST names",
                "names containing delisting marker 退",
                "N/C recent-listing markers",
                "codes outside the frozen prefix/board set",
            ],
            "note": "Prospective missed-opportunity audit denominator. Historical runs without an exact frozen universe remain exploratory.",
        },
        "summary": {
            "officialRawCount": len(all_items),
            "includedCount": len(included),
            "excludedCount": len(excluded),
            "exclusionCounts": counts,
        },
        "included": included,
        "excluded": excluded,
    }
    if len(included) < 3000:
        raise RuntimeError(f"included universe unexpectedly small: {len(included)}")

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(universe, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(universe["summary"], ensure_ascii=False))
    print(output)

if __name__ == "__main__":
    main()
